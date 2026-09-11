
import streamlit as st
st.markdown("""
    <style>
    #MainMenu {visibility: hidden;}
    footer {visibility: hidden;}
    header {visibility: hidden;}
    </style>
""",
            unsafe_allow_html=True)
import qrcode
from PIL import Image
import io
import os
import fitz  # PyMuPDF барои PDF
from docx import Document  # Барои Word
import openpyxl  # Барои Excel
from datetime import datetime
import uuid

# Папкаи асосӣ
BASE_UPLOAD_DIR = "levakant_multiusers_documents"
if not os.path.exists(BASE_UPLOAD_DIR):
    os.makedirs(BASE_UPLOAD_DIR)
7
# Рӯйхати давраҳои ҳисоботӣ
REPORT_PERIODS = ["3-моҳа", "6-моҳа", "9-моҳа", "Солона"]

# Рӯйхати муассисаҳо ва админ
USERS_DB = {
    "admin": {"password": "admin2026", "name": "Сардори Раёсат / Админ", "role": "admin"},
    "mtmu1": {"password": "123", "name": "МТМУ №1", "role": "user"},
    "mtmu2": {"password": "123", "name": "МТМУ №2", "role": "user"},
    "mtmu3": {"password": "123", "name": "МТМУ №3", "role": "user"},
    "mtmu4": {"password": "123", "name": "МТМУ №4", "role": "user"},
    "mtmu5": {"password": "123", "name": "МТМУ №5", "role": "user"},
    "mtmu6": {"password": "123", "name": "МТМУ №6", "role": "user"},
    "mtmu7": {"password": "123", "name": "МТМУ №7", "role": "user"},
    "mtmu8": {"password": "123", "name": "МТМУ №8", "role": "user"},
    "mtmu9": {"password": "123", "name": "МТМУ №9", "role": "user"},
    "mtmu10": {"password": "123", "name": "МТМУ №10", "role": "user"},
    "mtmu11": {"password": "123", "name": "МТМУ №11", "role": "user"},
    "mtmu12": {"password": "123", "name": "МТМУ №12", "role": "user"},
    "mtmu13": {"password": "123", "name": "МТМУ №13", "role": "user"},
    "mtmu14": {"password": "123", "name": "МТМУ №14", "role": "user"},
    "mtmu15": {"password": "123", "name": "МТМУ №15", "role": "user"},
    "litsey1": {"password": "123", "name": "Литсей №1", "role": "user"},
    "litsey2": {"password": "123", "name": "Литсей №2", "role": "user"},
}

# Ба таври худкор сохтани папкаҳои 3-моҳа, 6-моҳа, 9-моҳа ва Солона барои ҳар як муассиса
for u_code, u_info in USERS_DB.items():
    if u_info["role"] == "user":
        school_folder = os.path.join(BASE_UPLOAD_DIR, u_code)
        if not os.path.exists(school_folder):
            os.makedirs(school_folder)
        for period in REPORT_PERIODS:
            period_folder = os.path.join(school_folder, period)
            if not os.path.exists(period_folder):
                os.makedirs(period_folder)

# Танзимоти саҳифа
st.set_page_config(
    page_title="Вазорати маориф ва илми ҶТ - Бахши маорифи шаҳри Левакант",
    page_icon="📚",
    layout="wide"
)

# Тарҳи замонавӣ дар услуби портали давлатӣ
st.markdown("""
    <style>
    .stApp {
        background: #f1f5f9;
        color: #1e293b;
        font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif;
    }
    .top-navbar {
        background: #1b4d3e;
        padding: 12px 20px;
        color: white;
        display: flex;
        justify-content: space-between;
        align-items: center;
        border-radius: 8px;
        margin-bottom: 20px;
        box-shadow: 0 4px 12px rgba(0,0,0,0.1);
    }
    .top-navbar a {
        color: #cbd5e1;
        text-decoration: none;
        margin: 0 10px;
        font-size: 14px;
        font-weight: 500;
    }
    .header-box {
        background: linear-gradient(135deg, #1b4d3e 0%, #113328 100%);
        padding: 25px;
        border-radius: 12px;
        color: white;
        text-align: center;
        box-shadow: 0 10px 25px -5px rgba(27, 77, 62, 0.3);
        margin-bottom: 25px;
        border-bottom: 4px solid #c59b27;
    }
    .login-box {
        background: white;
        padding: 40px;
        border-radius: 16px;
        box-shadow: 0 10px 25px rgba(0,0,0,0.08);
        max-width: 400px;
        margin: auto;
        border-top: 5px solid #1b4d3e;
    }
    .stButton>button {
        background: #1b4d3e;
        color: white;
        border-radius: 8px;
        font-weight: 600;
        width: 100%;
        padding: 0.65rem;
        border: none;
        box-shadow: 0 4px 10px rgba(27, 77, 62, 0.2);
        transition: 0.3s;
    }
    .stButton>button:hover {
        background: #256b56;
    }
    .card-box {
        background: white;
        padding: 25px;
        border-radius: 12px;
        box-shadow: 0 4px 15px rgba(0,0,0,0.05);
        margin-bottom: 20px;
        border-left: 5px solid #1b4d3e;
    }
    </style>
""", unsafe_allow_html=True)

# Менюи болоӣ
st.markdown("""
    <div class="top-navbar">
        <div>
            <b>eHukumat</b> | Низоми электронии идоракунии маориф
        </div>
        <div>
            <a href="#">Оид ба мақомот</a>
            <a href="#">Фаъолият</a>
            <a href="#">Ҳуҷҷатҳо</a>
            <a href="#">Тамос</a>
        </div>
    </div>
""", unsafe_allow_html=True)

if "logged_in" not in st.session_state:
    st.session_state.logged_in = False
    st.session_state.username = ""
    st.session_state.role = ""
    st.session_state.user_fullname = ""

query_params = st.query_params
verify_id = query_params.get("doc_id", None)

# Қисми санҷиши ҳуҷҷат тавассути QR
if verify_id:
    st.markdown("""
        <div class="header-box">
            <h4 style="margin:0; font-size: 15px; color: #facc15;">ВАЗОРАТИ МАОРИФ ВА ИЛМИ ҶУМҲУРИИ ТОҶИКИСТОН</h4>
            <h2 style="margin:8px 0; font-size: 22px; color: #ffffff;">БАХШИ МАОРИФИ ШАҲРИ ЛЕВАКАНТ</h2>
            <p style="margin:0; color: #e2e8f0; font-size: 13px;">Системаи марказии тасдиқи ҳаққонияти ҳуҷҷатҳо</p>
        </div>
    """, unsafe_allow_html=True)

    found_file_path = None
    found_file_name = None

    for root, dirs, files in os.walk(BASE_UPLOAD_DIR):
        for file in files:
            if verify_id in file:
                found_file_name = file
                found_file_path = os.path.join(root, file)
                break
        if found_file_path:
            break

    st.markdown("<div class='card-box'>", unsafe_allow_html=True)
    if found_file_path:
        st.success("✅ САНҶИШИ МУВАФФАҚ: Ҳуҷҷати мазкур **РАСМӢ, АСЛӢ ва ЭЪТИБОРНОК** аст!")
        st.info(f"📋 Рақами ягонаи бақайдгирии ID: **{verify_id}**")
        st.write("Ин ҳуҷҷат аз ҷониби муассисаи марбута тасдиқ гардида, дар сервери давлатӣ ҳифз шудааст.")

        file_ext = found_file_name.split('.')[-1].lower()
        with open(found_file_path, "rb") as f:
            file_bytes = f.read()

        st.download_button(
            label=f"📥 Зеркашии нусхаи аслии ҳуҷҷат ({file_ext.upper()})",
            data=file_bytes,
            file_name=found_file_name,
            mime="application/octet-stream"
        )
    else:
        st.error("❌ ДИҚҚАТ: Ҳуҷҷат бо ин ID дар сервери расмӣ ёфт нашуд ё қалбакӣ аст!")
    st.markdown("</div>", unsafe_allow_html=True)

# Қисми вуруд ба система (Login)
elif not st.session_state.logged_in:
    st.markdown("<br>", unsafe_allow_html=True)
    col1, col2, col3 = st.columns([1, 2, 1])
    with col2:
        st.markdown("""
            <div class="login-box">
                <div style="text-align: center; margin-bottom: 20px;">
                    <span style="font-size: 40px;">🏛️</span>
                    <h4 style="color: #1b4d3e; margin-top: 10px; font-size: 16px;">Вуруд ба Низоми Левакант</h4>
                    <p style="color: #64748b; font-size: 12px;">Барои муассисаҳои таълимӣ</p>
                </div>
        """, unsafe_allow_html=True)

        username_input = st.text_input("Логин (масалан: mtmu1, litsey1 ё admin)")
        password_input = st.text_input("Парол", type="password")

        st.write("")
        if st.button("Ворид шудан"):
            if username_input in USERS_DB and USERS_DB[username_input]["password"] == password_input:
                st.session_state.logged_in = True
                st.session_state.username = username_input
                st.session_state.role = USERS_DB[username_input]["role"]
                st.session_state.user_fullname = USERS_DB[username_input]["name"]
                st.success("Хуш омадед!")
                st.rerun()
            else:
                st.error("❌ Логин ё парол хато аст!")

        st.markdown("""
            <div style='margin-top: 20px; font-size: 11px; color: #64748b; background: #f8fafc; padding: 10px; border-radius: 8px;'>
                <b>Маълумот:</b><code></code> аст.<br>
                 <code></code><code></code>)
            </div>
            </div>
        """, unsafe_allow_html=True)

# Қисми асосии коркард
else:
    col_top1, col_top2 = st.columns([6, 1])
    with col_top2:
        if st.button("🚪 Баромад"):
            st.session_state.logged_in = False
            st.session_state.username = ""
            st.session_state.role = ""
            st.rerun()

    st.markdown(f"""
        <div class="header-box">
            <h4 style="margin:0; font-size: 14px; color: #facc15;">ВАЗОРАТИ МАОРИФ ВА ИЛМИ ҶУМҲУРИИ ТОҶИКИСТОН</h4>
            <h1 style="margin:5px 0; font-size: 22px; color: #ffffff;">БАХШИ МАОРИФИ ШАҲРИ ЛЕВАКАНТ</h1>
            <p style="margin:0; color: #cbd5e1; font-size: 12px;">Муассисаи воридкунанда: <b>{st.session_state.user_fullname}</b></p>
        </div>
    """, unsafe_allow_html=True)

    if st.session_state.role == "user":
        st.markdown(f"### 🚀 Боргузорӣ ва танзими QR-код барои: {st.session_state.user_fullname}")

        # Интихоби давраи ҳисоботӣ
        report_period = st.selectbox(
            "📅 Папкаи давраи ҳисоботиро интихоб кунед:",
            options=REPORT_PERIODS
        )

        uploaded_file = st.file_uploader("Файли расмиро интихоб кунед (PDF, Word ё Excel)",
                                         type=["pdf", "docx", "xlsx"])

        if uploaded_file:
            st.success(f"✅ Файл қабул шуд: **{uploaded_file.name}**")

        doc_title = st.text_input("Номи Ҳуҷҷат / Фармоиш", placeholder="Фармоиши рақами...")
        signer_fio = st.text_input("Ф.И.О. Шахси масъул / Директор", placeholder="Холматов Б.")
        base_website = st.text_input("Линки Сервер (URL)", placeholder="https://your-app.streamlit.app")

        st.markdown("### 📍 Мавқеи ҷойгиршавии QR-код дар ҳуҷҷат")
        qr_position = st.selectbox(
            "Ҷойгиршавии QR-кодро интихоб кунед:",
            options=["Тарафи рост (Рости поён)", "Тарафи чап (Чапи поён)", "Дар байн (Маркази поён)"]
        )

        if st.button("✨ Тасдиқ ва сохтани QR-код"):
            if not uploaded_file or not doc_title or not signer_fio or not base_website:
                st.error("⚠️ Лутфан ҳамаи майдонҳоро пур кунед!")
            else:
                with st.spinner("Дар ҳоли коркард..."):
                    try:
                        unique_id = str(uuid.uuid4())[:8].upper()
                        now = datetime.now()
                        current_time_str = now.strftime("%Y-%m-%d %H:%M:%S")
                        folder_time_str = now.strftime("%Y-%m-%d_%H-%M-%S")

                        # Роҳ ба папкаи дахлдори муассиса ва давраи интихобшуда
                        user_folder = os.path.join(BASE_UPLOAD_DIR, st.session_state.username)
                        period_folder = os.path.join(user_folder, report_period)
                        time_folder = os.path.join(period_folder, folder_time_str)
                        if not os.path.exists(time_folder):
                            os.makedirs(time_folder)

                        file_extension = uploaded_file.name.split('.')[-1].lower()
                        saved_filename = f"{unique_id}_{uploaded_file.name}"
                        server_file_path = os.path.join(time_folder, saved_filename)

                        verification_url = f"{base_website.strip('/')}/?doc_id={unique_id}"
                        qr_data = f"=== ВАЗОРАТИ МАОРИФ ВА ИЛМИ ҶТ ===\nМУАССИСА: {st.session_state.user_fullname}\nДавра: {report_period}\nID: {unique_id}\nҲуҷҷат: {doc_title}\nМасъул: {signer_fio}\nСана: {current_time_str}\nСанҷиш: {verification_url}"

                        qr_img = qrcode.make(qr_data)
                        qr_byte_arr = io.BytesIO()
                        qr_img.save(qr_byte_arr, format='PNG')
                        qr_bytes = qr_byte_arr.getvalue()

                        if file_extension == 'pdf':
                            pdf_bytes = uploaded_file.read()
                            pdf_document = fitz.open(stream=pdf_bytes, filetype="pdf")
                            page = pdf_document[-1]
                            rect = page.rect

                            qr_size = 90
                            y1 = rect.height - 40
                            y0 = y1 - qr_size

                            if qr_position == "Тарафи чап (Чапи поён)":
                                x0 = 40
                                x1 = x0 + qr_size
                            elif qr_position == "Дар байн (Маркази поён)":
                                x0 = (rect.width - qr_size) / 2
                                x1 = x0 + qr_size
                            else:
                                x1 = rect.width - 40
                                x0 = x1 - qr_size

                            bbox = fitz.Rect(x0, y0, x1, y1)
                            page.insert_image(bbox, stream=qr_bytes)
                            output_data = pdf_document.write()
                            pdf_document.close()
                        elif file_extension == 'docx':
                            output_data = uploaded_file.read()
                            doc = Document(io.BytesIO(output_data))
                            doc.add_paragraph(
                                f"\n[{st.session_state.user_fullname} | Давра: {report_period} | Масъул: {signer_fio} | ID: {unique_id} | {current_time_str}]")
                            doc.add_picture(io.BytesIO(qr_bytes), width=100)
                            doc_io = io.BytesIO()
                            doc.save(doc_io)
                            output_data = doc_io.getvalue()
                        else:
                            output_data = uploaded_file.read()

                        with open(server_file_path, "wb") as f_out:
                            f_out.write(output_data)

                        st.success(
                            f"🎉 Ҳуҷҷат дар папкаи **{report_period}** бомуваффақият сабт шуд! ID: **{unique_id}**")
                        st.image(qr_img, width=130, caption=f"QR-код дар мавқеи: {qr_position}")

                        st.download_button(
                            label=f"📥 Зеркашии ҳуҷҷати имзошуда ({file_extension.upper()})",
                            data=output_data,
                            file_name=f"Imzoshuda_{unique_id}_{uploaded_file.name}",
                            mime="application/octet-stream"
                        )
                    except Exception as e:
                        st.error(f"Хатогӣ: {e}")

        st.markdown("---")
        st.markdown(f"### 📂 Ҳуҷҷатҳои боргузоришудаи {st.session_state.user_fullname}:")

        user_folder = os.path.join(BASE_UPLOAD_DIR, st.session_state.username)
        if os.path.exists(user_folder):
            my_files_found = False
            for root, dirs, files in os.walk(user_folder):
                for f_name in files:
                    my_files_found = True
                    full_p = os.path.join(root, f_name)
                    rel_p = os.path.relpath(full_p, user_folder)
                    st.write(f"📄 Папка ва файл: `{rel_p}`")
            if not my_files_found:
                st.info("Шумо ҳанӯз ягон ҳуҷҷат нафиристодаед.")

    elif st.session_state.role == "admin":
        st.markdown("### 🛡️ Панели Назорати Админ (Поисковик ва Бойгонии Давраҳо)")

        if os.path.exists(BASE_UPLOAD_DIR):
            # Поисковик
            search_query = st.text_input("🔍 Ҷустуҷӯи фаврӣ аз рӯи номи ҳуҷҷат:",
                                         placeholder="Номи ҳуҷҷат ё қисми онро нависед...")
            st.markdown("---")

            total_files = 0
            for root, dirs, files in os.walk(BASE_UPLOAD_DIR):
                total_files += len(files)

            st.info(f"Ҷамъи умумии ҳуҷҷатҳо дар сервер: {total_files} адад")

            if search_query:
                st.markdown("#### 🔎 Натиҷаҳои ҷустуҷӯ:")
                matched_files = []
                for root, dirs, files in os.walk(BASE_UPLOAD_DIR):
                    for file in files:
                        if search_query.lower() in file.lower():
                            matched_files.append((root, file))

                if matched_files:
                    for root_path, file_name in matched_files:
                        rel_path = os.path.relpath(root_path, BASE_UPLOAD_DIR)
                        path_parts = rel_path.split(os.sep)
                        school_code = path_parts[0] if len(path_parts) > 0 else "N/A"
                        period_name = path_parts[1] if len(path_parts) > 1 else "N/A"
                        time_dir = path_parts[2] if len(path_parts) > 2 else "N/A"

                        school_name = USERS_DB.get(school_code, {}).get("name", school_code)

                        col_a, col_b, col_c = st.columns([3, 2, 1])
                        with col_a:
                            st.write(f"📄 **{file_name}**")
                            st.caption(f"🏛️ Муассиса: {school_name} | 📅 Давра: {period_name}")
                        with col_b:
                            st.write(f"🕒 Сана: `{time_dir}`")
                        with col_c:
                            file_path = os.path.join(root_path, file_name)
                            with open(file_path, "rb") as f_dl:
                                st.download_button(
                                    label="📥 Зеркашӣ",
                                    data=f_dl.read(),
                                    file_name=file_name,
                                    key=f"search_{file_path}"
                                )
                        st.markdown("---")
                else:
                    st.warning("⚠️ Ҳуҷҷате бо ин ном ёфт нашуд.")

            # Намоиши сохтори папкаҳо мувофиқи муассиса ва папкаҳои ҳатмии давраҳо (3-моҳа, 6-моҳа ва ғ.)
            st.markdown("#### 📁 Феҳристи муассисаҳо ва папкаҳои ҳисоботӣ:")
            for school_code, u_info in USERS_DB.items():
                if u_info["role"] == "user":
                    school_path = os.path.join(BASE_UPLOAD_DIR, school_code)
                    school_name = u_info["name"]

                    with st.expander(f"📁 Муассиса: {school_name} ({school_code})"):
                        for period in REPORT_PERIODS:
                            period_path = os.path.join(school_path, period)
                            st.markdown(f"📂 **Давраи ҳисоботӣ: {period}**")

                            if os.path.exists(period_path) and os.listdir(period_path):
                                time_dirs = os.listdir(period_path)
                                for time_dir in sorted(time_dirs, reverse=True):
                                    time_path = os.path.join(period_path, time_dir)
                                    if os.path.isdir(time_path):
                                        st.markdown(f"&nbsp;&nbsp;&nbsp;&nbsp; 🕒 Вақт: `{time_dir}`")
                                        files_in_time = os.listdir(time_path)
                                        for file in files_in_time:
                                            file_path = os.path.join(time_path, file)
                                            col_a, col_b = st.columns([4, 1])
                                            with col_a:
                                                st.write(f"&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp; 📄 {file}")
                                            with col_b:
                                                with open(file_path, "rb") as f_dl:
                                                    st.download_button(
                                                        label="📥 Зеркашӣ",
                                                        data=f_dl.read(),
                                                        file_name=file,
                                                        key=f"folder_{file_path}"
                                                    )
                            else:
                                st.markdown("&nbsp;&nbsp;&nbsp;&nbsp; <i>(Папка холӣ аст)</i>", unsafe_allow_html=True)
                            st.markdown("---")
        else:
            st.warning("Ҳоло дар сервер ягон папка мавҷуд нест.")
